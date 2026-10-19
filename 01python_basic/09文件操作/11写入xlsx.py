# 引入openpyxl
import openpyxl
import random

# 1.创建一个wb
wb = openpyxl.Workbook()


# 2.创建一个sheet对象
# sheet = wb.create_sheet("期末考试成绩66")


sheet = wb.active
sheet.title = "期末考试77"

# 写入单元格标题
# 假设 姓名 语文 数学 英语
# sheet["A1"] = "姓名"
# sheet["B1"] = "语文"
# sheet["C1"] = "数学"
# sheet["D1"] = "英语"

# sheet.cell(row=行号, column=列号,value=值)



# 循环写入标题
title = ["姓名","语文","数学","英语"]
for i, v in enumerate(title):
    # 给第一行的1 2 3 4列加内容
    sheet.cell(row=1, column=i+1, value=v)

# 循环写入第一列的姓名
names = ["小田","小白","小盖","小岳","小 牛"]
for i, v in enumerate(names):
    # 给第二行开始 往后的每一行的第一列写入姓名
    sheet.cell(i + 2, 1, v)
    # 按列循环三次 写入成绩
    for j in range(1, 4):
        # 给每一行的第2 3 4列写入成绩
        sheet.cell(i + 2, j + 1, random.randint(0, 100))



# 保存数据
wb.save("utils/学生期末成绩.xls")
print("写入结束")



