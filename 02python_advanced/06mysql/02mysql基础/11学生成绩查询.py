# 查询数据
# 需求 查询到指定的班级的学生
from utilsmysql import fetch_data
def select_student():
    sql = "select name,sex,age,class_id from students where class_id = %s"
    params = (5,)
    fetch_data(sql,params)


# 三表联合查询 将最终的结果导出到一个excel表格中 表头是中文 (姓名 课程 成绩)
import openpyxl
import pandas as pd
def three_select_table():
    sql = "SELECT st.name AS '学生姓名',c.name AS '课程名',sc.score AS '分数' FROM students st INNER JOIN scores sc ON st.id = sc.student_id INNER JOIN courses c ON c.id = sc.course_id"
    result = fetch_data(sql)
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "学生成绩表"

    title = ["学生姓名","课程名","成绩"]
    for i,v in enumerate(title):
        sheet.cell(row=1,column=i+1,value=v)

    for r, row in enumerate(result, start=2):
        for c, key in enumerate(["学生姓名", "课程名", "分数"], start=1):
            sheet.cell(row=r, column=c, value=row[key])
    wb.save("utils/test.xlsx")
    print("end")

def getData(page,size):
    # page 页码 size 每页条码数 （1，2）
    state = (page - 1) * size # 分页公式
    """
        page
        1 跳过0条拿5条出来
        2 跳过5条拿5条
        3 跳过10条拿5条
        n (n-1)*size
    """
    sql = """select s.name AS 姓名, c.name AS 科目,sc.score AS 成绩
            from students s
            join scores sc ON s.id = sc.student_id
            join courses c ON c.id = sc.course_id
            ORDER BY sc.id
            limit %s,%s
    """
    data = fetch_data(sql, (state, size))
    print(data)

def exportDate():
    sql = """select s.name AS 姓名, c.name AS 科目,sc.score AS 成绩
            from students s
            join scores sc ON s.id = sc.student_id
            join courses c ON c.id = sc.course_id
            ORDER BY sc.id
    """
    data = fetch_data(sql)
    df = pd.DataFrame(data)
    df.to_excel("utils/test.xlsx",index=False)


if __name__ == '__main__':
    # select_student()
    # three_select_table()
    # getData(2,7)
    exportDate()